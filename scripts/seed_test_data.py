"""Seed deterministic test datasets for connector smoke tests.

Generates five small datasets — customers, products, orders, employees,
transactions — in CSV / XLSX / JSON / Parquet under ``seed-data/<format>/``
and (optionally) loads them into Postgres + MinIO.

Designed to support manual exercises of the connector P0 stack
(CSV / Excel / JSON / Parquet / S3 / PostgreSQL). Production seeding goes
through Alembic; this script only writes to a dedicated ``seed`` schema
that is dropped & recreated on each run.

Usage
-----
    python scripts/seed_test_data.py           # generate + postgres + minio
    python scripts/seed_test_data.py --no-postgres
    python scripts/seed_test_data.py --no-minio
    python scripts/seed_test_data.py --only-generate

Environment variables (with defaults shown):

    DATABASE_URL              postgresql://postgres:postgres@localhost:5436/dataquality_db
    SEED_PG_SCHEMA            seed
    SEED_MINIO_ENDPOINT       localhost:9000
    SEED_MINIO_ACCESS_KEY     minioadmin
    SEED_MINIO_SECRET_KEY     <falls back to MINIO_ROOT_PASSWORD>
    SEED_MINIO_BUCKET         dq-data-assets
    SEED_MINIO_SECURE         false
"""
from __future__ import annotations

import argparse
import json
import logging
import os
import random
import sys
from dataclasses import dataclass
from datetime import datetime, timedelta, timezone
from decimal import Decimal
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional

logger = logging.getLogger("seed_test_data")

REPO_ROOT = Path(__file__).resolve().parents[1]
SEED_ROOT = REPO_ROOT / "seed-data"
FORMATS: tuple[str, ...] = ("csv", "xlsx", "json", "parquet")

# Deterministic seed → reproducible builds.
RNG = random.Random(20260427)


# ─────────────────────────────────────────────────────────────────────────
# Dataset generators
# ─────────────────────────────────────────────────────────────────────────


@dataclass(frozen=True)
class Dataset:
    name: str
    rows: List[Dict[str, Any]]

    @property
    def columns(self) -> List[str]:
        return list(self.rows[0].keys()) if self.rows else []


def _iso(dt: datetime) -> str:
    return dt.replace(microsecond=0).isoformat()


def build_customers(n: int = 50) -> Dataset:
    cities = ["Paris", "Berlin", "Madrid", "Rome", "Amsterdam", "Lisbon", "Vienna"]
    rows: List[Dict[str, Any]] = []
    for i in range(1, n + 1):
        rows.append(
            {
                "customer_id": i,
                "first_name": f"First{i:03d}",
                "last_name": f"Last{i:03d}",
                "email": f"customer{i:03d}@example.com" if i % 11 else None,
                "city": RNG.choice(cities),
                "signup_at": _iso(
                    datetime(2024, 1, 1, tzinfo=timezone.utc)
                    + timedelta(days=RNG.randint(0, 700))
                ),
                "is_active": i % 7 != 0,
            }
        )
    return Dataset("customers", rows)


def build_products(n: int = 30) -> Dataset:
    categories = ["books", "electronics", "apparel", "home", "outdoor"]
    rows = []
    for i in range(1, n + 1):
        rows.append(
            {
                "product_id": i,
                "sku": f"SKU-{i:04d}",
                "name": f"Product {i:03d}",
                "category": RNG.choice(categories),
                "unit_price": round(RNG.uniform(5.0, 499.99), 2),
                "in_stock": RNG.randint(0, 250),
            }
        )
    return Dataset("products", rows)


def build_orders(customers: Dataset, products: Dataset, n: int = 200) -> Dataset:
    statuses = ["pending", "paid", "shipped", "cancelled"]
    rows = []
    for i in range(1, n + 1):
        cust = RNG.choice(customers.rows)
        prod = RNG.choice(products.rows)
        qty = RNG.randint(1, 6)
        rows.append(
            {
                "order_id": i,
                "customer_id": cust["customer_id"],
                "product_id": prod["product_id"],
                "quantity": qty,
                "total_amount": round(prod["unit_price"] * qty, 2),
                "status": RNG.choice(statuses),
                "ordered_at": _iso(
                    datetime(2025, 1, 1, tzinfo=timezone.utc)
                    + timedelta(days=RNG.randint(0, 450))
                ),
            }
        )
    return Dataset("orders", rows)


def build_employees(n: int = 25) -> Dataset:
    departments = ["engineering", "sales", "support", "operations", "finance"]
    rows = []
    for i in range(1, n + 1):
        manager = RNG.randint(1, i - 1) if i > 5 and RNG.random() < 0.7 else None
        rows.append(
            {
                "employee_id": i,
                "full_name": f"Employee {i:03d}",
                "department": RNG.choice(departments),
                "manager_id": manager,
                "hire_date": _iso(
                    datetime(2018, 1, 1, tzinfo=timezone.utc)
                    + timedelta(days=RNG.randint(0, 2500))
                ),
                "salary": round(RNG.uniform(45_000, 165_000), 2),
                "metadata": {"level": RNG.randint(1, 7), "remote": RNG.random() > 0.4},
            }
        )
    return Dataset("employees", rows)


def build_transactions(orders: Dataset, n: int = 300) -> Dataset:
    currencies = ["EUR", "USD", "GBP", "CHF"]
    rows = []
    for i in range(1, n + 1):
        order = RNG.choice(orders.rows)
        rows.append(
            {
                "transaction_id": i,
                "order_id": order["order_id"],
                "currency": RNG.choice(currencies),
                "amount": float(
                    Decimal(str(order["total_amount"]))
                    * Decimal(str(round(RNG.uniform(0.95, 1.05), 4)))
                ),
                "processor": RNG.choice(["stripe", "adyen", "paypal"]),
                "occurred_at": _iso(
                    datetime(2025, 6, 1, tzinfo=timezone.utc)
                    + timedelta(minutes=RNG.randint(0, 60 * 24 * 365))
                ),
            }
        )
    return Dataset("transactions", rows)


def build_all() -> List[Dataset]:
    customers = build_customers()
    products = build_products()
    orders = build_orders(customers, products)
    employees = build_employees()
    transactions = build_transactions(orders)
    return [customers, products, orders, employees, transactions]


# ─────────────────────────────────────────────────────────────────────────
# File writers
# ─────────────────────────────────────────────────────────────────────────


def _ensure_dirs() -> None:
    for fmt in FORMATS:
        (SEED_ROOT / fmt).mkdir(parents=True, exist_ok=True)


def write_csv(ds: Dataset) -> Path:
    import csv

    target = SEED_ROOT / "csv" / f"{ds.name}.csv"
    with target.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=ds.columns)
        writer.writeheader()
        for row in ds.rows:
            writer.writerow(
                {k: (json.dumps(v) if isinstance(v, dict) else v) for k, v in row.items()}
            )
    return target


def write_json(ds: Dataset) -> Path:
    target = SEED_ROOT / "json" / f"{ds.name}.json"
    target.write_text(
        json.dumps(ds.rows, indent=2, default=str), encoding="utf-8"
    )
    return target


def write_xlsx(ds: Dataset) -> Path:
    from openpyxl import Workbook

    target = SEED_ROOT / "xlsx" / f"{ds.name}.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = ds.name[:31]
    ws.append(ds.columns)
    for row in ds.rows:
        ws.append(
            [
                json.dumps(v) if isinstance(v, dict) else v
                for v in row.values()
            ]
        )
    wb.save(target)
    wb.close()
    return target


def write_parquet(ds: Dataset) -> Path:
    import pyarrow as pa  # type: ignore
    import pyarrow.parquet as pq  # type: ignore

    target = SEED_ROOT / "parquet" / f"{ds.name}.parquet"
    rows = [
        {
            k: (json.dumps(v) if isinstance(v, dict) else v)
            for k, v in row.items()
        }
        for row in ds.rows
    ]
    table = pa.Table.from_pylist(rows)
    pq.write_table(table, target)
    return target


def write_all_formats(datasets: List[Dataset]) -> List[Path]:
    _ensure_dirs()
    written: List[Path] = []
    for ds in datasets:
        written.append(write_csv(ds))
        written.append(write_json(ds))
        written.append(write_xlsx(ds))
        written.append(write_parquet(ds))
    return written


# ─────────────────────────────────────────────────────────────────────────
# Postgres loader
# ─────────────────────────────────────────────────────────────────────────


_PG_TYPE_MAP = {
    int: "BIGINT",
    float: "DOUBLE PRECISION",
    bool: "BOOLEAN",
    str: "TEXT",
}


def _infer_pg_type(values: Iterable[Any]) -> str:
    """Pick a permissive Postgres type from sampled values."""
    seen: set[type] = set()
    for v in values:
        if v is None:
            continue
        if isinstance(v, dict):
            return "JSONB"
        seen.add(type(v))
    if not seen:
        return "TEXT"
    if seen == {int}:
        return "BIGINT"
    if seen <= {int, float}:
        return "DOUBLE PRECISION"
    if seen == {bool}:
        return "BOOLEAN"
    return "TEXT"


def load_postgres(datasets: List[Dataset], *, schema: str, dsn: str) -> None:
    import psycopg2
    import psycopg2.extras

    logger.info("postgres_load: connecting dsn=%s schema=%s", _redact(dsn), schema)
    conn = psycopg2.connect(dsn)
    conn.autocommit = True
    try:
        with conn.cursor() as cur:
            cur.execute(f'DROP SCHEMA IF EXISTS "{schema}" CASCADE')
            cur.execute(f'CREATE SCHEMA "{schema}"')

            for ds in datasets:
                col_types = {
                    col: _infer_pg_type(row.get(col) for row in ds.rows)
                    for col in ds.columns
                }
                col_def = ", ".join(
                    f'"{col}" {col_types[col]}' for col in ds.columns
                )
                cur.execute(
                    f'CREATE TABLE "{schema}"."{ds.name}" ({col_def})'
                )

                rows = []
                for row in ds.rows:
                    rows.append(
                        tuple(
                            json.dumps(v) if isinstance(v, dict) else v
                            for v in (row[c] for c in ds.columns)
                        )
                    )
                placeholders = ", ".join(["%s"] * len(ds.columns))
                psycopg2.extras.execute_batch(
                    cur,
                    f'INSERT INTO "{schema}"."{ds.name}" '
                    f'({", ".join(chr(34) + c + chr(34) for c in ds.columns)}) '
                    f"VALUES ({placeholders})",
                    rows,
                    page_size=200,
                )
                logger.info(
                    "postgres_load: %s.%s rows=%d", schema, ds.name, len(rows)
                )
    finally:
        conn.close()


def _redact(dsn: str) -> str:
    # Hide password component from log output.
    if "://" not in dsn or "@" not in dsn:
        return dsn
    head, tail = dsn.split("://", 1)
    if "@" not in tail:
        return dsn
    creds, host = tail.split("@", 1)
    if ":" in creds:
        user = creds.split(":", 1)[0]
        return f"{head}://{user}:***@{host}"
    return dsn


# ─────────────────────────────────────────────────────────────────────────
# MinIO loader
# ─────────────────────────────────────────────────────────────────────────


def load_minio(
    files: List[Path],
    *,
    endpoint: str,
    access_key: str,
    secret_key: str,
    bucket: str,
    secure: bool,
) -> None:
    from minio import Minio  # type: ignore

    logger.info(
        "minio_load: endpoint=%s bucket=%s secure=%s", endpoint, bucket, secure
    )
    client = Minio(
        endpoint, access_key=access_key, secret_key=secret_key, secure=secure
    )
    if not client.bucket_exists(bucket):
        client.make_bucket(bucket)
        logger.info("minio_load: created bucket=%s", bucket)

    for path in files:
        # Use ``<format>/<name>.<ext>`` as the object key.
        key = f"{path.parent.name}/{path.name}"
        client.fput_object(bucket, key, str(path))
        logger.info("minio_load: %s -> s3://%s/%s", path.name, bucket, key)


# ─────────────────────────────────────────────────────────────────────────
# CLI
# ─────────────────────────────────────────────────────────────────────────


def main(argv: Optional[List[str]] = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    parser.add_argument(
        "--only-generate",
        action="store_true",
        help="Only write files; skip Postgres + MinIO loaders.",
    )
    parser.add_argument(
        "--no-postgres",
        action="store_true",
        help="Skip the Postgres loader step.",
    )
    parser.add_argument(
        "--no-minio",
        action="store_true",
        help="Skip the MinIO upload step.",
    )
    parser.add_argument(
        "--log-level", default=os.environ.get("SEED_LOG_LEVEL", "INFO")
    )
    args = parser.parse_args(argv)

    logging.basicConfig(
        level=getattr(logging, args.log_level.upper(), logging.INFO),
        format="%(asctime)s %(levelname)s %(name)s :: %(message)s",
    )

    datasets = build_all()
    written = write_all_formats(datasets)
    logger.info("generate: wrote %d files under %s", len(written), SEED_ROOT)

    if args.only_generate:
        return 0

    failures = 0

    if not args.no_postgres:
        try:
            dsn = os.environ.get(
                "DATABASE_URL",
                "postgresql://postgres:postgres@localhost:5436/dataquality_db",
            )
            schema = os.environ.get("SEED_PG_SCHEMA", "seed")
            load_postgres(datasets, schema=schema, dsn=dsn)
        except Exception as exc:  # pragma: no cover - integration path
            failures += 1
            logger.error("postgres_load_failed: %s", exc)

    if not args.no_minio:
        try:
            secret = os.environ.get("SEED_MINIO_SECRET_KEY") or os.environ.get(
                "MINIO_ROOT_PASSWORD"
            )
            if not secret:
                raise RuntimeError(
                    "Missing MinIO secret. Set SEED_MINIO_SECRET_KEY or "
                    "MINIO_ROOT_PASSWORD before running the seeder."
                )
            load_minio(
                written,
                endpoint=os.environ.get("SEED_MINIO_ENDPOINT", "localhost:9000"),
                access_key=os.environ.get(
                    "SEED_MINIO_ACCESS_KEY",
                    os.environ.get("MINIO_ROOT_USER", "minioadmin"),
                ),
                secret_key=secret,
                bucket=os.environ.get("SEED_MINIO_BUCKET", "dq-data-assets"),
                secure=os.environ.get("SEED_MINIO_SECURE", "false").lower()
                == "true",
            )
        except Exception as exc:  # pragma: no cover - integration path
            failures += 1
            logger.error("minio_load_failed: %s", exc)

    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
